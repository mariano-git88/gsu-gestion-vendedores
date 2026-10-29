"""
Tests de la marca de primer pedido BIZ en el lector de pedidos.

Los pedidos de la plataforma BIZ llegan con dos artículos "fantasma" de
$0,01: `MKT CARTA BIZ UY` (vino por BIZ) y `MKT BIZ 1ER PED UY` (es el
primero, corresponde 5%). En oct-2026 dos primeros pedidos se cargaron sin
el 5% y hubo que corregirlos a mano; la app los marca a partir de esto.

Son puros: arman un xlsx con el layout de la plantilla, sin red.
"""

import io
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pedidos  # noqa: E402


def _xlsx(*hojas):
    """Cada hoja: lista de (codigo, cantidad, precio)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for n, items in enumerate(hojas, 1):
        ws = wb.create_sheet(f"Pedido {n}")
        ws["C3"], ws["E3"] = "cliente de prueba", 9999
        ws["B7"] = "ADHESIVOS DE CONTACTO"  # encabezado de rubro
        for fila, (cod, cant, precio) in enumerate(items, start=8):
            ws.cell(fila, 2, cod)
            ws.cell(fila, 5, precio)
            ws.cell(fila, 6, cant)
            ws.cell(fila, 7, cant * precio)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_primer_pedido_biz_se_detecta():
    [p] = pedidos.leer_pedidos(_xlsx([
        ("MKT BIZ 1ER PED UY", 1, 0.01),
        ("MKT CARTA BIZ UY", 1, 0.01),
        ("PBD 125 M E", 15, 283.62),
    ]))
    assert p.es_primer_pedido_biz
    assert p.es_pedido_biz


def test_codigo_con_espacios_o_minusculas_igual_se_detecta():
    [p] = pedidos.leer_pedidos(_xlsx([
        ("  mkt biz  1er ped uy ", 1, 0.01),
        ("PBD 125 M E", 15, 283.62),
    ]))
    assert p.es_primer_pedido_biz


def test_pedido_biz_que_no_es_el_primero_no_lleva_descuento():
    [p] = pedidos.leer_pedidos(_xlsx([
        ("MKT CARTA BIZ UY", 1, 0.01),
        ("PBD 125 M E", 15, 283.62),
    ]))
    assert p.es_pedido_biz
    assert not p.es_primer_pedido_biz


def test_pedido_del_vendedor_no_se_marca():
    [p] = pedidos.leer_pedidos(_xlsx([("PBD 125 M E", 15, 283.62)]))
    assert not p.es_pedido_biz
    assert not p.es_primer_pedido_biz


def test_fantasma_con_cantidad_cero_no_cuenta():
    # Si algún día la plantilla trae el código precargado en el catálogo,
    # solo cuenta cuando tiene cantidad, como cualquier otro ítem.
    [p] = pedidos.leer_pedidos(_xlsx([
        ("MKT BIZ 1ER PED UY", 0, 0.01),
        ("PBD 125 M E", 15, 283.62),
    ]))
    assert not p.es_primer_pedido_biz


def test_cada_hoja_se_marca_por_separado():
    p1, p2 = pedidos.leer_pedidos(_xlsx(
        [("MKT BIZ 1ER PED UY", 1, 0.01), ("PBD 125 M E", 15, 283.62)],
        [("PBD 125 M E", 15, 283.62)],
    ))
    assert p1.es_primer_pedido_biz
    assert not p2.es_primer_pedido_biz
